"""Create a concise follow-up report for Claude's Phase 4 instructions."""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from project_paths import REPORTS_DIR

OUT = REPORTS_DIR / "TaniAdapt_Claude_Followup_Phase4.docx"
NAVY, PALE, GRID = "17365D", "F3F6F9", "D9D9D9"

def shade(cell, color):
    e = OxmlElement("w:shd"); e.set(qn("w:fill"), color); cell._tc.get_or_add_tcPr().append(e)

def border(cell):
    p = cell._tc.get_or_add_tcPr(); b = p.first_child_found_in("w:tcBorders")
    if b is None: b = OxmlElement("w:tcBorders"); p.append(b)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x = b.find(qn(f"w:{side}"))
        if x is None: x = OxmlElement(f"w:{side}"); b.append(x)
        x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "4"); x.set(qn("w:color"), GRID)

def put(cell, text, bold=False, white=False, size=8.3, center=False):
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(text)); r.bold = bold; r.font.name = "Aptos"; r.font.size = Pt(size)
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos"); r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    r.font.color.rgb = RGBColor(255,255,255) if white else RGBColor(0,0,0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER; border(cell)

def table(doc, heads, rows, widths):
    t = doc.add_table(rows=1, cols=len(heads)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(heads): shade(t.rows[0].cells[i], NAVY); put(t.rows[0].cells[i], h, True, True, center=True); t.rows[0].cells[i].width = Inches(widths[i])
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            if ri%2: shade(cells[i], PALE)
            put(cells[i], v, center=(i==0 or (i==1 and len(str(v)) < 18))); cells[i].width=Inches(widths[i])
        t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    doc.add_paragraph()

def heading(doc, s):
    p=doc.add_heading(s, level=1); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(4)
    for r in p.runs: r.font.color.rgb=RGBColor(0,0,0); r.font.name="Aptos Display"

def para(doc,s):
    p=doc.add_paragraph(s); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.08

def main():
    d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
    d.styles["Normal"].font.name="Aptos"; d.styles["Normal"].font.size=Pt(10)
    title=d.add_paragraph(style="Title"); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    tp=title._p.get_or_add_pPr(); b=OxmlElement("w:pBdr"); bot=OxmlElement("w:bottom"); bot.set(qn("w:val"),"nil"); b.append(bot); tp.append(b)
    title.add_run("TaniAdapt Phase 4 Follow Up for Claude")
    sub=d.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; sub.add_run("Implementation response to the requested audit corrections").italic=True
    para(d,"This report records what was changed after Claude's review instructions. It distinguishes completed corrections from remaining data or environment blockers. The project remains at Phase 4 close-candidate; Phase 5 remains draft pending ratification.")

    heading(d,"1 Instruction status")
    table(d,["No", "Status", "Implementation and evidence"],[
        ["1","PARTIAL","Clean venv was created but dependency installation did not persist, so clean rerun failed with missing NumPy. The canonical runner was completed with local project Python 3.11. See reports/logs/rerun_diff.md."],
        ["2","DONE","CHIRPS N corrected from 972 to 936 from raw matched data. 2016, 2020, and 2024 each contribute N=312. Indramayu has 24 months, Depok 12. Pooled r uses N=936; per-unit median uses 26 valid units."],
        ["3","DONE","WLS, Huber, and MAD-z sensitivity results now record bootstrap SE, CI, MDE, and inference method. All use district-cluster Webb bootstrap with B=999; Huber was recomputed."],
        ["4","DONE","All BPS source rows now have diisi_oleh=BELUM DIVERIFIKASI MANUSIA and verified=false. DS03 is PENDING_MANUAL_VERIFICATION in the registry."],
        ["5","DONE","Crop selection changed from FINAL RESMI to DRAFT. MDE padi is +/-0.0255; crop classes, robust N, unique weather cells, and horticulture output restrictions were revised."],
        ["6","DONE","Schema examples use placeholders or nulls. Unsupported thresholds and numerical example claims were removed. Metadata, percentile convention, evidence-source actions, and evidence_insufficient handling were added."],
        ["7","DONE","Local commit 365214c and immutable new tag phase4-close-candidate were created. Cleanup commit 227bd79 removed obsolete patchers and improved .gitignore."],
    ],[.4,.8,5.65])

    heading(d,"2 Reproduction result")
    para(d,"The clean-environment result is not accepted as reproduced. The virtual environment was created from requirements.txt, but dependency installation did not complete and the first runner step failed because NumPy was unavailable. The data and scripts were then rerun using the local Python 3.11 environment that already contains the statistical dependencies. No CDS credential, API key, or raw data content was printed, stored, or committed.")
    table(d,["Runner component","Result"],[
        ["A2 CHIRPS versus AgERA5", "Completed; raw matched output contains 936 unit-months."],
        ["A3 spatial cells and PCA", "Completed; 22 unique AgERA5 cells reported."],
        ["A4 panel inference and MDE", "Completed; main maize N=124 before the separately documented robust crop-selection screen."],
        ["Canonical result refresh", "Completed; DS09 and cross-dataset JSON updated."],
    ],[2.3,4.55])

    heading(d,"3 Corrected numerical record")
    table(d,["Metric","Prior","Current","Interpretation"],[
        ["CHIRPS matched rows","972","936","Raw-data count; no missing months are imputed."],
        ["CHIRPS r pooled","0.7102","0.7102, N=936","Estimate unchanged; denominator corrected."],
        ["CHIRPS r within year","r 0.6350 / 0.7277 / 0.6930","Same; N=312 each","2016 / 2020 / 2024 respectively."],
        ["CHIRPS unit median r","0.7569","0.7569, 26 units","Only units with a valid correlation contribute."],
        ["Maize main MDE","older audit 0.2306","0.1444","2.8 x Webb bootstrap SE 0.0516."],
        ["Rice MDE","older audit 0.0329","0.0255","2.8 x Webb bootstrap SE 0.0091; T=3 is not high power."],
        ["Huber CI / MDE","[-0.0292, 0.0709] / 0.0765","[-0.0282, 0.0704] / 0.0725","Re-estimated with Webb cluster bootstrap, B=999."],
    ],[1.3,1.5,1.55,2.5])

    heading(d,"4 Statistical guardrails now applied")
    table(d,["Model","Method now stored","Current result"],[
        ["Main maize","Webb wild-cluster bootstrap; cluster=kabupaten; B=999","beta +0.0368; CI [-0.0521, +0.1300]; MDE 0.1444."],
        ["WLS","WLS fixed-effect residualization; Webb bootstrap; B=999","beta +0.0395; SE 0.0307; MDE 0.0859."],
        ["Huber","Huber RLM; Webb wild-cluster residual bootstrap; B=999","beta +0.0198; SE 0.0259; CI [-0.0282, +0.0704]; MDE 0.0725."],
        ["MAD-z","Within-kabupaten empirical P95 screen; Webb bootstrap; B=999; 7 dropped","beta +0.0555; SE 0.0463; MDE 0.1298."],
    ],[1.05,3.15,2.65])
    para(d,"All confidence intervals include zero. These results are sensitivity estimates, not evidence of a causal weather effect, calibrated yield rule, or ML target.")

    heading(d,"5 Audit and documentation changes")
    table(d,["Area","Current status"],[
        ["CHIRPS polygon sampling","PENDING. The previous fixed reduction factor was removed because the repository has no actual polygon-area extraction. No before/after bias or range is claimed."],
        ["Rice BPS baseline","PENDING_MANUAL_VERIFICATION. Values were not filled or guessed. A human must inspect BPS Jabar Dalam Angka 2016, 2017, and 2018 PDFs."],
        ["Crop selection","DRAFT. Padi: mechanistic advisory. Jagung/kedelai: empirical sensitivity analysis plus mechanistic advisory. Cabai/tomat: Tier B literature-transfer rule. Bawang merah: blocked pending new data."],
        ["Horticulture output","Restricted to jumlah hari kondusif dalam jendela N hari terakhir. No 0-100 score, disease probability, disease occurrence, or yield-loss claim."],
        ["Engine schema","No literal thresholds or unsupported sensitivity numbers. It includes issue time, timezone, versions, freshness, rank_basis_n, percentile reporting convention, and evidence-source action status."],
    ],[1.7,5.15])

    heading(d,"6 Remaining decisions and blockers")
    table(d,["Item","Why unresolved","Required next step"],[
        ["Clean reproduction","A newly created venv lacks installed packages after installation attempt.","Complete pip installation in a stable network-enabled environment, then rerun scripts/run_phase4_pipeline.py."],
        ["Rice baseline","Cited BPS page entries are not human-verified.","Inspect the three primary PDFs; set verified=true only after checking the page and table."],
        ["Polygon sampling","No raw polygon-area weather extraction exists in the repository.","Acquire or construct real area-weighted CHIRPS and AgERA5 extracts, then calculate the comparison."],
        ["Phase 5","Selection document is deliberately DRAFT.","Ratify, revise, or narrow the crop portfolio and its permitted claims."],
    ],[1.45,2.65,2.75])

    heading(d,"7 Version record")
    para(d,"Phase 4 close-candidate commit: 365214cb2394f9f1ec4340be5bdf665c8a702227. Tag: phase4-close-candidate. The tag was created without moving any existing tag. Later repository cleanup commit: 227bd79874c7cdaaa03ad75ea274c456da3d5013.")
    para(d,"Primary files for review: reports/logs/rerun_diff.md; reports/results/09_climate_covariates_and_chirps.json; reports/results/10_cross_dataset_synthesis.json; reports/dataset_audit/final_dataset_registry.csv; docs/crop_outcome_decision.md; docs/engine_output_schema.md.")
    d.save(OUT); print(OUT)

if __name__ == "__main__": main()
