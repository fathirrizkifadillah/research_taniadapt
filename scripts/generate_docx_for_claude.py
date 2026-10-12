"""Generate a comprehensive, beautifully styled progress and audit report for Claude."""
from __future__ import annotations

import os
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from project_paths import PROJECT_ROOT

OUT_DIR = PROJECT_ROOT / "context" / "to_claude"
OUT_PATH = OUT_DIR / "TaniAdapt_Research_Progress_and_Audit_Closure_Report_for_Claude.docx"

NAVY = "1B365D"
STEEL = "2C5282"
PALE_BG = "F4F6F9"
BORDER_GRAY = "CCCCCC"
DARK_GRAY = "333333"
GREEN_TEXT = "1B6E34"
AMBER_TEXT = "B45309"
RED_TEXT = "B91C1C"


def set_cell_background(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), BORDER_GRAY)
        borders.append(b)
    tc_pr.append(borders)


def format_cell(cell, text: str, bold: bool = False, white: bool = False,
                size: float = 8.5, center: bool = False, color_override: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")

    if white:
        run.font.color.rgb = RGBColor(255, 255, 255)
    elif color_override:
        r_hex = int(color_override[0:2], 16)
        g_hex = int(color_override[2:4], 16)
        b_hex = int(color_override[4:6], 16)
        run.font.color.rgb = RGBColor(r_hex, g_hex, b_hex)
    else:
        run.font.color.rgb = RGBColor(40, 40, 40)

    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_borders(cell)


def add_table(doc, headers: list[str], rows: list[list[str]], col_widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        set_cell_background(cell, NAVY)
        format_cell(cell, h, bold=True, white=True, size=8.8, center=True)
        cell.width = Inches(col_widths[i])
    hdr_row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    for r_idx, r_data in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        bg = PALE_BG if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            if bg != "FFFFFF":
                set_cell_background(cell, bg)
            is_center = (c_idx == 0 or (c_idx == 1 and len(str(val)) <= 15))
            color_override = None
            if str(val).startswith("DONE") or str(val) == "PASS" or str(val) == "VERIFIED":
                color_override = GREEN_TEXT
            elif str(val).startswith("PARTIAL") or str(val) == "DRAFT" or str(val) == "PENDING":
                color_override = AMBER_TEXT
            elif str(val).startswith("TIDAK BISA") or str(val) == "CANNOT_OPEN" or str(val) == "UNVERIFIED":
                color_override = RED_TEXT

            format_cell(cell, val, bold=(c_idx == 0), white=False, size=8.3, center=is_center, color_override=color_override)
            cell.width = Inches(col_widths[c_idx])

    doc.add_paragraph()


def add_callout(doc, title: str, text: str, border_color: str = NAVY, bg_color: str = PALE_BG) -> None:
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(7.0)
    set_cell_background(cell, bg_color)

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left_b = OxmlElement("w:left")
    left_b.set(qn("w:val"), "single")
    left_b.set(qn("w:sz"), "24")
    left_b.set(qn("w:space"), "0")
    left_b.set(qn("w:color"), border_color)
    borders.append(left_b)

    for side in ("top", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "none")
        borders.append(b)
    tc_pr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}\n")
    r_title.bold = True
    r_title.font.name = "Aptos"
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor(27, 54, 93)

    r_text = p.add_run(text)
    r_text.font.name = "Aptos"
    r_text.font.size = Pt(8.8)
    r_text.font.color.rgb = RGBColor(50, 50, 50)
    doc.add_paragraph()


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()

    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)

    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(9.5)

    # Title Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t_run = title_p.add_run("TANIADAPT — RESEARCH PROGRESS & AUDIT REPORT")
    t_run.bold = True
    t_run.font.name = "Aptos Display"
    t_run.font.size = Pt(17)
    t_run.font.color.rgb = RGBColor(27, 54, 93)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(8)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s_run = sub_p.add_run("Evidence-Based Agricultural Decision Support System for West Java Smallholders\nComprehensive Research Progress Dossier & Claude Instruction Audit Response")
    s_run.font.name = "Aptos"
    s_run.font.size = Pt(10)
    s_run.font.color.rgb = RGBColor(80, 80, 80)
    s_run.italic = True

    add_callout(
        doc,
        "CORE ARCHITECTURAL PRINCIPLE & METHODOLOGICAL CONTRACT",
        "Pipeline: Evidence → Finding → Agronomic Interpretation → Model → Decision Logic → Action → LLM Explanation.\n"
        "Strict Safeguards: Dilarang membuat ambang/threshold sendiri; dilarang menggunakan kata 'membuktikan'; Decision Engine didecoupling dari LLM; sistem wajib mengembalikan status 'evidence_insufficient' saat parameter tidak didukung data terverifikasi.",
        border_color=NAVY,
        bg_color=PALE_BG,
    )

    # Section 1: Executive Overview
    h1 = doc.add_heading("1. Executive Research Overview & Positioning", level=1)
    h1.paragraph_format.space_before = Pt(8); h1.paragraph_format.space_after = Pt(3)
    for r in h1.runs: r.font.name = "Aptos Display"; r.font.color.rgb = RGBColor(27, 54, 93)

    p1 = doc.add_paragraph(
        "TaniAdapt adalah platform Precision Agro-Advisory dan Agricultural Decision Support System (DSS) untuk "
        "petani Jawa Barat. Berbeda dengan aplikasi generik yang hanya menyajikan prakiraan cuaca atau chatbot LLM "
        "yang menebak diagnosis secara arbitrer, TaniAdapt memisahkan secara tegas antara:\n"
        "1. Decision Engine Numerik/Mekanistik: Otak ilmiah deterministik berbasis neraca air FAO-56 Penman-Monteith, "
        "koefisien tanaman Kc per fase, dan proksi mikroklimat persentil lokal historis.\n"
        "2. LLM + RAG Layer: Lapisan komunikasi bahasa alami yang menerjemahkan output terstruktur engine ke bahasa petani "
        "dan menjelaskan alasan agronomi (decision trace) tanpa menciptakan ambang sendiri."
    )
    p1.paragraph_format.space_after = Pt(4); p1.paragraph_format.line_spacing = 1.1

    add_table(
        doc,
        ["Phase", "Deskripsi Fase Riset", "Status Saat Ini", "Capaian & Output Utama"],
        [
            ["Phase 0", "Definisi Masalah Pertanian Jabar", "DONE", "Celah keputusan petani: air, pupuk, kalender tanam, risiko OPT."],
            ["Phase 1", "Research Questions & Hypotheses", "DONE", "Pra-registrasi hipotesis fenologi & ekonometrika di docs/hypotheses.md."],
            ["Phase 2", "Deep Literature Review", "DONE", "Ekstraksi 16 poin dari 16 paper internasional/nasional di context/."],
            ["Phase 3", "Dataset Discovery & Fusion", "DONE", "Akuisisi 9 dataset (DS01–DS09): AgERA5, CHIRPS, BPS, Distanhor."],
            ["Phase 4", "Dataset Quality Audit", "CLOSED (Candidate)", "Resolusi 11 blocker A1–A11, eliminasi artefak data, audit 22 sel AgERA5."],
            ["Phase 5", "Crop & Outcome Selection", "DRAFT (Pending)", "Portfolio 5 tanaman: Padi, Jagung, Kedelai, Cabai, Tomat di crop_outcome_decision.md."],
            ["Phase 6", "Exploratory Data Analysis (EDA)", "DONE", "10 Jupyter Notebooks terverifikasi (0 errors) + figur sintesis cross-dataset."],
            ["Phase 7", "Statistical Analysis & Panel FE", "READY TO ENGAGE", "Wild cluster bootstrap Webb (B=999), MDE analitis, Two-Way FE."],
            ["Phase 8", "Temporal Feature Engineering", "READY TO ENGAGE", "Rolling 14-day precip, CDD, water balance (P - ET0), P10 dew depression."],
            ["Phase 9", "ML Baseline & Sensitivity", "READY TO ENGAGE", "Uji sensitivitas WLS, Huber RLM, MAD-z (bukan penaksir kausal semu)."],
            ["Phase 10", "Evidence → Risk Model", "IN PROGRESS", "Penyusunan evidence registry resmi (FAO-56 ET0, Kc Table 12, P10 IFWC)."],
            ["Phase 11", "Risk → Decision Engine", "IN PROGRESS", "Validasi modul mekanistik FAO-56 Penman-Monteith & skema output JSON."],
            ["Phase 12", "Non-Yield Validation", "VALIDATED", "Hindcast 2016 basah, 2019 kering, 2023 kering (Hit rate: 88.9%–100%)."],
            ["Phase 13", "TaniAdapt Web Application", "NEXT MILESTONE", "Integrasi antarmuka pengguna, visualisasi cuaca, & lapisan eksplanasi LLM."]
        ],
        [0.8, 1.9, 1.4, 2.9]
    )

    # Section 2: Claude Prompt A Implementation
    h2 = doc.add_heading("2. Response & Audit Resolution for Claude's Prompt A (Tasks 1–7)", level=1)
    h2.paragraph_format.space_before = Pt(8); h2.paragraph_format.space_after = Pt(3)
    for r in h2.runs: r.font.name = "Aptos Display"; r.font.color.rgb = RGBColor(27, 54, 93)

    add_table(
        doc,
        ["No", "Tugas Claude (Prompt A)", "Status", "Rincian Implementasi & Hasil Bukti Empiris"],
        [
            ["1", "REPRODUKSI Pipeline Phase 4", "DONE (Local Rerun)",
             "scripts/run_phase4_pipeline.py dijalankan dan seluruh tahapan A2, A3, A4, serta refresh hasil kanonik lolos 100%. "
             "Perbedaan lingkungan bersih (venv gagal pasang numpy) vs Python 3.11 lokal dicatat jujur di reports/logs/rerun_diff.md."],
            ["2", "Koreksi CHIRPS N & Polygon Claim", "DONE",
             "N unit-bulan diperbaiki menjadi 936 (312 per tahun untuk 2016, 2020, 2024; Indramayu 24, Depok 12). "
             "Korelasi pooled r=0.7102 (N=936), median per-unit r=0.7569 (26 unit). Kata 'membuktikan' pada sampling poligon dihapus "
             "dan status diturunkan ke PENDING karena data poligon area mentah belum diekstraksi."],
            ["3", "SE Robust Jagung & Bounds", "DONE",
             "Model sensitivitas H_CORN_01 (WLS, Huber, MAD-z) dihitung ulang menggunakan Wild Cluster Bootstrap Webb (cluster=kabupaten, B=999). "
             "MDE dihitung baku dari 2.8 x SE bootstrap: Main MDE=0.1444, WLS MDE=0.0859, Huber MDE=0.0725, MAD-z MDE=0.1298. "
             "Frasa 'arah bervariasi' dihapus, seluruh CI mencakup nol dilaporkan apa adanya di reports/results/10_cross_dataset_synthesis.json."],
            ["4", "BPS Manual Input & Registry Flag", "DONE",
             "reports/dataset_audit/bps_manual_input.csv diisi dengan diisi_oleh='BELUM DIVERIFIKASI MANUSIA' dan sumber_halaman_verified=false. "
             "Baseline padi 2015-2017 resmi dikunci berstatus PENDING_MANUAL_VERIFICATION di docs/phase4_final_dataset_evidence_registry.md."],
            ["5", "Penyempurnaan docs/crop_outcome_decision.md", "DONE",
             "Status diubah ke DRAFT pending ratifikasi. MDE padi konsisten +/-0.0255 log points. Sel grid unik dinyatakan 18 sentra (provinsi cabai/tomat "
             "tanpa sel terdefinisi). Klaim ingatan Garut >60% dihapus. Registry classes diselaraskan. Skor 0-100 diganti jumlah hari kondusif. "
             "Filter N_valid memakai robust outlier flag."],
            ["6", "Pembersihan docs/engine_output_schema.md", "DONE",
             "Seluruh angka estimasi palsu pada contoh diganti '<dari hasil engine>' atau null. Padi berstatus published_model (2015-2017). "
             "Jagung dilarang mengklaim status 'stres' saat CI mencakup nol. Cabai/tomat berstatus ambang: PENDING_LITERATURE. "
             "Action IDs disinkronkan identik dengan tabel aksi; aksi fungisida diberi status PENDING_EXTENSION_SOURCE. Aturan persentil didefinisikan."],
            ["7", "Pembaruan Git History & Tagging", "DONE",
             "reports/logs/git_history.md diperbarui. Tag git baru 'phase4-close-candidate' telah dibuat di commit 365214cb2394f9f1ec4340be5bdf665c8a702227 "
             "tanpa memindahkan tag git lama. Commit lanjutan perbaikan disimpan di 0a54f50."]
        ],
        [0.4, 1.8, 1.2, 3.6]
    )

    # Section 3: Claude Prompt B Implementation
    h3 = doc.add_heading("3. Response & Implementation for Claude's Prompt B (Tasks 1–5)", level=1)
    h3.paragraph_format.space_before = Pt(8); h3.paragraph_format.space_after = Pt(3)
    for r in h3.runs: r.font.name = "Aptos Display"; r.font.color.rgb = RGBColor(27, 54, 93)

    add_table(
        doc,
        ["Tugas", "Komponen (Prompt B)", "Status", "Temuan & Realisasi Teknis"],
        [
            ["Tugas 1", "Klimatologi Baseline AgERA5 1991–2020", "PARTIAL / DONE (Sel Tersedia)",
             "175 file tahunan (untuk 6 sel unik) berhasil diunduh via CDS API ke data/raw/DS01_agera5_climatology_1991_2020/. "
             "Sisa 16 sel terkendala antrean CDS API (~4 jam waktu proses). Skrip scripts/compute_11_climatology_baseline.py "
             "telah menghitung persentil P10-P90 per jendela 14-hari & musiman MH/MK. Uji konsistensi mencatat pergeseran suhu median "
             "2015-2024 lebih hangat +0.2C s.d. +0.4C dibanding 1991-2020. Output: reports/results/11_climatology_baseline.json & PNG."],
            ["Tugas 2", "Evidence Registry Literasi Agronomi", "DONE",
             "scripts/build_evidence_registry.py menghasilkan reports/dataset_audit/evidence_registry.csv dan docs/evidence_registry_literature.md. "
             "Entri VERIFIED resmi: FAO56_ET0_PM (Eq. 6), FAO56_KC_RICE, FAO56_KC_MAIZE, FAO56_KC_SOY, FAO56_KC_TOMATO (FAO-56 Table 12, p. 110). "
             "Entri UNVERIFIED/CANNOT_OPEN: Suhu kardinal per fase tanaman, air cabai tropis, cuaca kondusif antraknosa/hawar daun, serta SOP aksi penyuluhan."],
            ["Tugas 3", "Parameter Mekanistik & Validasi FAO-56", "DONE",
             "KATAM berstatus CANNOT_ACCESS (dokumen konsep terbuka, API kalender per kecamatan tidak tersedia terbuka). "
             "Implementasi Python mandiri FAO-56 Penman-Monteith di scripts/mechanistic/fao56.py divalidasi terhadap Contoh 18 Bab 4 FAO-56 "
             "(Ekspektasi: 3.88 mm/hari vs Dihitung: 3.877 mm/hari -> PASS). Disimpan di reports/results/12_mechanistic_params.json."],
            ["Tugas 4", "Hindcast Validasi Non-Yield (2016, 2019, 2023)", "VALIDATED",
             "Pra-registrasi peristiwa ekstrem dikunci di docs/hindcast_plan.md pada commit c515e27 SEBELUM eksekusi. "
             "scripts/evaluate_hindcast.py mengevaluasi data harian AgERA5 2015-2024 di 18 sentra: "
             "HC_2016_WET (La Nina basah Okt-Des 2016) = 18/18 HIT (100.0%); "
             "HC_2019_DRY (Kemarau IOD Jun-Okt 2019) = 18/18 HIT (100.0%); "
             "HC_2023_DRY (El Nino Jun-Nov 2023) = 16/18 HIT, 2 MISS (88.9%). Output: reports/results/13_hindcast_validation.json."],
            ["Tugas 5", "Penyelarasan Dokumen Seleksi & Skema Output", "DONE",
             "docs/crop_outcome_decision.md dan docs/engine_output_schema.md telah disinkronkan hanya mengutip ID bukti berstatus VERIFIED "
             "(FAO56_ET0_PM, FAO56_KC_RICE, FAO56_KC_MAIZE, FAO56_KC_SOY, FAO56_KC_TOMATO)."]
        ],
        [0.8, 1.8, 1.3, 3.1]
    )

    # Section 4: Evidence Registry Audit Summary
    h4 = doc.add_heading("4. Evidence Registry Audit: Verified vs Unverified Citations", level=1)
    h4.paragraph_format.space_before = Pt(8); h4.paragraph_format.space_after = Pt(3)
    for r in h4.runs: r.font.name = "Aptos Display"; r.font.color.rgb = RGBColor(27, 54, 93)

    add_table(
        doc,
        ["Evidence ID", "Komoditas", "Variabel / Parameter", "Nilai Terverifikasi", "Status Verifikasi", "Sitasi & Halaman Dokumen Sumber"],
        [
            ["FAO56_ET0_PM", "Semua", "Reference Evapotranspiration (ET0)", "Persamaan Penman-Monteith (Eq. 6)", "VERIFIED", "Allen et al. (1998), FAO-56 Chapter 2/Eq. 6, p. 65."],
            ["FAO56_KC_RICE", "Padi", "Koefisien Tanaman (Kc)", "Kc_ini=1.05, mid=1.20, end=0.90-0.60", "VERIFIED", "FAO-56 Table 12, p. 110 (Subhumid standard)."],
            ["FAO56_KC_MAIZE", "Jagung", "Koefisien Tanaman (Kc)", "Kc_ini=0.30, mid=1.20, end=0.60-0.35", "VERIFIED", "FAO-56 Table 12, p. 110 (Field corn)."],
            ["FAO56_KC_SOY", "Kedelai", "Koefisien Tanaman (Kc)", "Kc_ini=0.40, mid=1.15, end=0.50", "VERIFIED", "FAO-56 Table 12, p. 110 (Soybeans)."],
            ["FAO56_KC_TOMATO", "Tomat", "Koefisien Tanaman (Kc)", "Kc_ini=0.60, mid=1.15, end=0.70-0.90", "VERIFIED", "FAO-56 Table 12, p. 110 (Tomato)."],
            ["RICE_TEMP_PHASES", "Padi", "Suhu Kardinal (Base/Opt/Max)", "PENDING (Tidak disubstitusi)", "UNVERIFIED", "Tidak ada paper varietas lokal terbuka pada pass ini."],
            ["MAIZE_TEMP_PHASES", "Jagung", "Suhu Kardinal (Base/Opt/Max)", "PENDING (Tidak disubstitusi)", "UNVERIFIED", "Tidak ada paper varietas lokal terbuka pada pass ini."],
            ["SOY_TEMP_PHASES", "Kedelai", "Suhu Kardinal (Base/Opt/Max)", "PENDING (Tidak disubstitusi)", "UNVERIFIED", "Tidak ada paper varietas lokal terbuka pada pass ini."],
            ["CHILI_WATER", "Cabai", "Koefisien Tanaman (Kc)", "PENDING (Sweet pepper tidak diinterpolasi)", "UNVERIFIED", "FAO-56 Table 12 opened but not representative for cabai."],
            ["CHILI_ANTHRACNOSE", "Cabai", "Proksi Cuaca Antraknosa", "PENDING (Ambang literal dilarang)", "CANNOT_OPEN", "Belum ada paper formula proksi kebasahan operasional terverifikasi."],
            ["TOMATO_LATE_BLIGHT", "Tomat", "Proksi Hawar Daun Phytophthora", "PENDING (Ambang literal dilarang)", "CANNOT_OPEN", "Belum ada paper formula proksi kebasahan operasional terverifikasi."],
            ["ACT_DRAINAGE_CHECK", "Semua", "Aksi Teknis Drainase", "PENDING_EXTENSION_SOURCE", "CANNOT_OPEN", "Menunggu pembukaan dokumen SOP resmi BPTP/Balitsa."],
            ["ACT_FUNGICIDE_REV", "Horti", "Aksi Evaluasi Fungisida", "PENDING_EXTENSION_SOURCE", "CANNOT_OPEN", "Menunggu pembukaan dokumen SOP resmi Kementan/Balitsa."]
        ],
        [1.3, 0.6, 1.4, 1.5, 0.9, 1.3]
    )

    # Section 5: Unresolved Blockers and Next Steps
    h5 = doc.add_heading("5. Unresolved Data Limitations & Next Steps (Phase 7–13)", level=1)
    h5.paragraph_format.space_before = Pt(8); h5.paragraph_format.space_after = Pt(3)
    for r in h5.runs: r.font.name = "Aptos Display"; r.font.color.rgb = RGBColor(27, 54, 93)

    p_limit = doc.add_paragraph(
        "Batasan Data yang Tidak Dapat Diselesaikan (Unresolved Data Limitations):\n"
        "1. Ketiadaan Data BPS KSA Padi Level Kabupaten 2018–2020: Padi dibatasi pada baseline terverifikasi 2015–2017.\n"
        "2. Zero Temporal Overlap Sayuran Kabupaten: Hortikultura (Cabai & Tomat) dikunci permanen di Tier B (Rule-based mikroklimat persentil P10), dilarang menggunakan regresi yield atau supervised ML penyakit.\n"
        "3. Ketiadaan Sensor Kebasahan Daun Fisik: Ambang depresiasi embun (Tmin - Td) fajar dikonvensikan sebagai proksi persentil relatif lokal P10, bukan bukti basah daun fisik.\n"
        "4. Unduh Penuh AgERA5 1991–2020: 175 file (6 sel) terunduh; 16 sel sisa tertunda oleh rate-limit antrean CDS API (~4 jam).\n\n"
        "Langkah Menuju Phase 7–13:\n"
        "• Phase 7 & 9: Mengunci pemodelan panel ekonometrika pada batas MDE (bukan klaim elastisitas kausal).\n"
        "• Phase 10 & 11: Membangun Decision Engine berbasis modul mekanistik terverifikasi (FAO-56 ET0 + Kc) dan aturan persentil relatif.\n"
        "• Phase 13: Menghubungkan output terstruktur engine ke antarmuka aplikasi TaniAdapt dan lapisan eksplanasi LLM/RAG."
    )
    p_limit.paragraph_format.line_spacing = 1.1

    doc.save(OUT_PATH)
    print(f"Successfully generated Claude report at: {OUT_PATH}")


if __name__ == "__main__":
    main()
