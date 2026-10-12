"""Create a review dossier for an external research reviewer."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from project_paths import PROJECT_ROOT, REPORTS_DIR


OUTPUT = REPORTS_DIR / "TaniAdapt_Progress_and_Evidence_Pack_for_Claude.docx"
NAVY = "17365D"
BLUE = "DDEBF7"
PALE = "F3F6F9"
GRID = "D9D9D9"


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), GRID)


def set_cell_text(cell, text: str, *, bold: bool = False, white: bool = False,
                  size: float = 8.5, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.color.rgb = RGBColor(255, 255, 255) if white else RGBColor(0, 0, 0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell_border(cell)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, NAVY)
        set_cell_text(cell, header, bold=True, white=True, size=8.5,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        if widths:
            cell.width = Inches(widths[i])
    header_props = table.rows[0]._tr.get_or_add_trPr()
    header_props.append(OxmlElement("w:tblHeader"))
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if idx % 2:
                shade(cells[i], PALE)
            align = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 1) and len(str(value)) < 30 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], str(value), size=8.1, align=align)
            if widths:
                cells[i].width = Inches(widths[i])
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.name = "Aptos Display" if level == 1 else "Aptos"
    return p


def paragraph(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("TaniAdapt research dossier | Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)
    run.font.size = Pt(8)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)
    add_page_number(sec)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10)

    title = doc.add_paragraph(style="Title")
    title_props = title._p.get_or_add_pPr()
    title_border = title_props.find(qn("w:pBdr"))
    if title_border is not None:
        title_props.remove(title_border)
    title_border = OxmlElement("w:pBdr")
    no_bottom_border = OxmlElement("w:bottom")
    no_bottom_border.set(qn("w:val"), "nil")
    title_border.append(no_bottom_border)
    title_props.append(title_border)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("TaniAdapt Research Progress and Evidence Pack")
    r.font.name = "Aptos Display"
    r.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("For independent review and direction setting by Claude").italic = True
    subtitle.paragraph_format.space_after = Pt(16)

    paragraph(doc, "Purpose. This dossier consolidates the research status, dataset audit, executed exploratory data analysis, evidence limits, and unresolved decisions in the TaniAdapt repository. It is written so the reviewer can decide the next research direction without assuming that a prior project decision is final.", "Purpose.")
    paragraph(doc, "Review request. Assess whether the evidence is sufficient to formally close Phase 4, whether the proposed crop and outcome portfolio is defensible, and which analysis path should continue into Phases 5 to 9.", "Review request.")

    heading(doc, "1 Project scope and research logic")
    paragraph(doc, "TaniAdapt is a research-first agricultural decision-support system for smallholders in West Java. Its intended product is an evidence-bounded advisory system that translates weather, crop phenology, and agronomic evidence into explainable actions. It is not yet a production application or a validated predictive model.")
    add_bullets(doc, [
        "A deterministic computation and rule engine must produce the structured risk output.",
        "A language model may narrate that output, but must not create thresholds, probabilities, yield estimates, or causal claims.",
        "Each advisory rule needs an evidence tier and an explicit list of claims that are allowed and forbidden.",
        "When data do not support an inference, the engine should return evidence_insufficient rather than make a stronger claim.",
    ])

    heading(doc, "2 Phase progress")
    add_table(doc,
        ["Phase", "Status", "Progress and evidence"],
        [
            ["0 Problem definition", "Complete", "Problem framing and user context defined."],
            ["1 Research questions and hypotheses", "Complete", "Questions and hypotheses established before data work."],
            ["2 Literature review", "Complete", "Literature evidence collected for crop and climate mechanisms."],
            ["3 Dataset discovery", "Complete", "Nine primary datasets plus cross-dataset synthesis identified."],
            ["4 Dataset quality audit", "Closure ready", "Audit outputs, registry, deprecation rules, and reproducibility runner prepared. Pipeline has not been rerun in the current environment."],
            ["5 Crop and outcome selection", "Decision draft exists", "A selection document exists, but the reviewer should decide whether to ratify, revise, or postpone it."],
            ["6 EDA", "Executed exploratory work", "Ten notebooks exist with stored executed outputs. EDA supports audit conclusions; it is not a final confirmatory analysis."],
            ["7 Statistical analysis", "Partial", "Cross-dataset panel sensitivity and MDE work completed; no causal claims should be made."],
            ["8 Feature engineering", "Not started", "Requires a ratified crop, outcome, period, and evidence tier."],
            ["9 ML baseline", "Not started", "Blocked for supervised outcome paths until targets and validation design are defensible."],
            ["10 to 13 Risk model to application", "Not started", "Architecture is drafted, but implementation should follow evidence decisions."],
        ], [1.25, 1.15, 4.5])

    heading(doc, "3 Dataset inventory and audit disposition")
    paragraph(doc, "The table below uses the new canonical registry. Its purpose is to define what each dataset can support, rather than to rank datasets by apparent completeness.")
    add_table(doc,
        ["ID", "Dataset grain and period", "Audit disposition", "Permitted research role"],
        [
            ["DS01", "Daily AgERA5; 27 ADM2 mapped to 22 unique cells; 2015-2024", "Approved with limits", "Weather backbone and local relative climatology"],
            ["DS02", "Bangladesh district by crop-season by year; 1,728 rows; 2015-2023", "Approved benchmark", "Panel-regression methodology benchmark only"],
            ["DS03", "West Java ADM2 annual rice productivity; raw 2015-2020", "Conditional", "Reconciled rice baseline only for 2015-2017; mechanistic path"],
            ["DS04", "Daily 0.5 degree terrestrial grid; 12 West Java points; 2016-2025", "Internal only", "Define extreme-climate indicators and internal consistency checks"],
            ["DS05", "ADM2 annual maize productivity; 18 active regencies", "Conditional", "Sensitivity analysis and mechanistic maize advisory"],
            ["DS06", "Provincial annual horticulture productivity plus weather", "Tier B only", "Rule-based IFWC microclimate advisory for chili and tomato"],
            ["DS07", "ADM2 annual soybean productivity and geometry; 17 sentra", "Conditional", "Secondary panel benchmark and mechanistic soybean advisory"],
            ["DS08", "ADM2 production and harvest area", "Supporting", "Outcome reconciliation and production decomposition"],
            ["DS09", "Daily CHIRPS by ADM2 plus monthly ONI and DMI", "Approved with limits", "Macroclimate context and partial precipitation benchmark"],
        ], [0.45, 2.25, 1.15, 3.05])

    heading(doc, "4 Audit results and hard limits")
    add_table(doc,
        ["Topic", "Audit finding", "Research implication"],
        [
            ["Spatial dependence", "27 administrative areas share 22 AgERA5 weather cells; five pairs share a cell.", "Do not treat 27 ADM2 weather series as independent observations. Spatial dependence and clustering matter."],
            ["Rice outcome", "2015-2017 area-weighted productivity agrees with BPS. Years 2018-2020 are unverified; 2019 shows a systematic reporting artifact of about +19% to +22%.", "Do not use rice as the primary weather-yield panel. Use a mechanistic seasonal advisory until verified outcomes are obtained."],
            ["Maize outcome", "In 2022, 13 regencies carry identical 75.70 ku/ha values, consistent with uniform imputation.", "Exclude 2022 from the main panel. Main period is 2015-2021 with 124 active observations in 18 regencies."],
            ["Soybean outcome", "2022 spatial SD collapses to 0.19 ku/ha versus 2.18 in 2021, indicating likely imputation.", "Exclude 2022. Use 2015-2021 and 17 sentra selected ex ante, yielding 108 valid observations."],
            ["Horticulture", "Only provincial outcome data overlap weather. ADM2 horticulture data are 2013-2014 and there are no disease or OPT labels.", "No annual yield regression, disease classifier, or calibrated disease probability for chili or tomato."],
            ["DS04 independence", "DS04 is ERA5 lineage and correlates strongly with DS01 (reported anomaly r = 0.9932).", "It is an internal consistency comparison, not independent validation."],
            ["CHIRPS coverage", "Independent benchmark data only exist for 2016, 2020, and 2024; 972 valid ADM2-month records.", "Do not describe this as 10-year validation and do not fill blocked years without a new acquisition route."],
        ], [1.05, 2.9, 2.95])

    heading(doc, "5 Executed exploratory data analysis")
    paragraph(doc, "Ten notebooks contain executed cells and stored outputs. The EDA was used to inspect data structure, temporal consistency, spatial dependence, descriptive variation, artifacts, and cross-dataset agreement. The items below summarize what was explored and what the exploratory results support.")
    eda_rows = [
        ["01 AgERA5", "Daily coverage; physical consistency; climatology; spatial PCA; residual variability after fixed effects", "27 ADM2 map to 22 cells; 100% physical consistency; demeaned PCA PC1 plus PC2 explains 80.97% of spatial variance; mean pairwise r = 0.6671."],
        ["02 Bangladesh rice", "Balance, season productivity, district and temporal variance decomposition, two-way FE benchmark", "Balanced 64 x 3 x 9 panel with 1,728 observations. Intended for method transfer only, not Indonesian parameter transfer."],
        ["03 West Java rice", "Area-weighted reconciliation to BPS; annual pattern; artifact check; within-district variation", "2015-2017 reconciliation difference is 0.01%, 0%, and 0%. 2019 artifact and T = 3 limit empirical panel inference."],
        ["04 Agroclimatic", "Grid coverage, CDD percentile distribution, macroclimate plausibility, DS01 anomaly comparison", "12 official-polygon points. High DS01 agreement reflects shared ERA5 lineage rather than independent validation."],
        ["05 Maize", "Panel completeness, spatial and time variance, 2022 duplicate-value anomaly, detrended Subang residuals", "2022 flagged suspect. Subang detrended residual CV is 5.49%; district and residual variation remain material after exclusion."],
        ["06 Horticulture", "Provincial productivity variation, temporal overlap, availability of disease labels", "Provincial CV is 4.49% for chili besar and 5.71% for chili rawit. No usable ADM2 outcome-weather overlap and no OPT labels."],
        ["07 Soybean and geography", "Ex-ante sentra selection, yearly spatial dispersion, centroid to weather-cell lookup", "17 active sentra selected. 2022 SD collapse flags imputation. ADM2-to-AgERA5 mean mapping distance is 5.2 km."],
        ["08 Production and area", "Yield identity check, area share, production-change decomposition, denominator sensitivity", "Yield identity is 100% consistent for active maize and soybean. Mean 67.16% of production change is associated with harvest-area change."],
        ["09 Climate covariates", "ENSO/IOD episodes, monthly precipitation anomalies, CHIRPS availability and bias", "ONI/DMI complete for 120 months. CHIRPS comparison is partial: within-year r = 0.635, 0.7277, 0.693 for 2016, 2020, 2024."],
        ["10 Cross dataset", "Panel matrix, feature distributions, weather benchmark, spatial sampling, robustness, MDE, variance decomposition", "Integrates the audit and sensitivity evidence. It does not establish a validated predictive or causal model."],
    ]
    add_table(doc, ["Notebook", "EDA performed", "Main result"], eda_rows, [1.15, 2.95, 2.8])

    heading(doc, "6 Cross dataset statistical results")
    paragraph(doc, "The panel work is sensitivity analysis. The available observational data do not identify causal weather effects because outcome quality, time depth, spatial dependence, and exposure measurement are constrained. Confidence intervals and minimum detectable effects are more informative than a binary significance interpretation.")
    add_table(doc,
        ["Quantity", "Canonical value", "Interpretation"],
        [
            ["Maize rainfall window estimate", "+0.0368 log points per SD", "Association only; not a causal effect."],
            ["Maize Webb wild-bootstrap 95% CI", "[-0.0521, +0.1300]", "Includes zero and small agronomically relevant effects."],
            ["Maize MDE", "+/- 0.1444 log points", "Based on 2.8 x bootstrap SE. Study cannot resolve smaller effects reliably."],
            ["Maize robust checks", "WLS +0.0395; Huber +0.0198; clean MAD-z +0.0555", "Direction and magnitude vary across reasonable specifications; do not turn into a calibrated rule."],
            ["Soybean MDE", "+/- 0.1250 log points", "Secondary benchmark with constrained power."],
            ["Rice MDE", "+/- 0.0255 log points", "Not evidence of high power; T = 3 depletes degrees of freedom."],
            ["AgERA5 CHIRPS pooled anomaly r", "0.7102", "Only across 36 valid benchmark months, not a full 10-year independent validation."],
            ["Maximum wet-season signed bias", "+72.31%", "Requires bias-aware communication where precipitation is high."],
        ], [2.0, 1.45, 3.45])

    heading(doc, "7 Existing crop and outcome decision draft")
    paragraph(doc, "A file named docs/crop_outcome_decision.md already proposes a portfolio. It should be treated as a draft for review because the stated project roadmap still identifies Phase 4 as current. The reviewer should decide whether the following choices should be formally ratified.")
    add_table(doc,
        ["Crop", "Proposed outcome and period", "Proposed evidence path", "Non negotiable limit"],
        [
            ["Rice", "ADM2 productivity, reconciled 2015-2017", "Mechanistic seasonal advisory: calendar, GDD, water balance", "No primary panel weather-yield regression until BPS outcome verification."],
            ["Maize", "ADM2 annual productivity, 2015-2021", "Sensitivity panel plus mechanistic phenology", "No causal or calibrated yield claim; exclude 2022."],
            ["Soybean", "ADM2 annual productivity, 2015-2021", "Secondary panel benchmark plus mechanistic advisory", "No use of 2022; only 17 ex-ante sentra."],
            ["Chili", "Provincial productivity, 2017-2024", "Tier B literature-transfer IFWC microclimate rule", "No disease classifier, probability, or yield regression."],
            ["Tomato", "Provincial productivity, 2017-2024", "Tier B shared IFWC microclimate rule", "No disease classifier, probability, or yield regression."],
        ], [0.7, 1.75, 2.25, 2.2])

    heading(doc, "8 Architecture status")
    paragraph(doc, "The architecture specification is conceptually aligned with the audit. It separates a seasonal layer for food crops and a daily layer for horticulture. The output contract requires crop, location, as-of date, indicators, status, evidence tier, evidence references, allowed claims, forbidden claims, limitations, and action identifiers.")
    add_bullets(doc, [
        "Seasonal layer: rice and maize are intended to use relative climatology, crop calendar, GDD, and water-balance indicators rather than unsupported season yield predictions.",
        "Daily layer: chili and tomato are intended to report accumulation of weather-favorable conditions, not calibrated probability of pest or disease occurrence.",
        "The schema should eventually add operational metadata such as data freshness, issue time, timezone, rule version, and data version.",
        "The schema's crop enum still contains bawang_merah while the decision draft locks a five-crop portfolio. This should be reconciled before implementation.",
    ])

    heading(doc, "9 Reproducibility and repository state")
    add_bullets(doc, [
        "A canonical CSV registry now records each dataset's grain, valid analytical period, status, permitted role, prohibited use, and result artifact.",
        "A Phase 4 registry document records canonical statistics and deprecates prior contradictory numbers: 25 versus 22 weather cells, 14 versus 12 DS04 points, broad CHIRPS validation claims, and older MDE values.",
        "Scripts now use scripts/project_paths.py rather than a hard-coded C drive root.",
        "requirements.txt and docs/reproducibility.md describe environment setup and the official Phase 4 runner.",
        "scripts/run_phase4_pipeline.py refreshes the A2, A3, and A4 canonical computations; it may optionally execute the ten notebooks.",
        "The runner has not been executed in the current environment after this cleanup. The evidence registry is therefore closure-ready, not independently revalidated by this dossier creation." ,
    ])

    heading(doc, "10 Decisions requested from Claude")
    paragraph(doc, "Please give an explicit decision for each item below. A useful response should distinguish what is approved for descriptive or mechanistic advisory use from what is approved for empirical inference, prediction, or a future ML baseline.")
    requests = [
        ["1", "Phase 4 closure", "Is the registry and audit sufficient to close Phase 4 after rerunning the canonical pipeline, or are additional audits required?"],
        ["2", "Crop portfolio", "Should the proposed five-crop portfolio be ratified, narrowed, or expanded?"],
        ["3", "Empirical outcome path", "Should maize and soybean remain sensitivity-only analyses, or should the project defer all empirical yield paths pending better outcome data?"],
        ["4", "Rice path", "Is a purely mechanistic rice advisory appropriate, and what minimum data would justify reopening its empirical path?"],
        ["5", "Horticulture path", "Is Tier B IFWC rule-based advisory an acceptable research output without disease labels, and what language or validation guardrails are necessary?"],
        ["6", "Phase 6 and 7 scope", "Which EDA findings should become confirmatory analyses, and what preregistered sensitivity plan is needed?"],
        ["7", "ML eligibility", "What data, target definition, validation split, and external validation threshold are required before Phase 9 can begin?"],
        ["8", "Architecture corrections", "Which schema gaps must be fixed before a deterministic engine is implemented?"],
    ]
    add_table(doc, ["No", "Decision", "Question for review"], requests, [0.4, 1.5, 4.95])

    heading(doc, "11 Source artifacts in the repository")
    paragraph(doc, "The following artifacts are the primary evidence base for this dossier. They should be inspected before changing any result or evidence tier.")
    add_bullets(doc, [
        "reports/results/01_AgERA5.json through reports/results/09_climate_covariates_and_chirps.json",
        "reports/results/10_cross_dataset_synthesis.json",
        "reports/dataset_audit/final_dataset_registry.csv",
        "docs/phase4_final_dataset_evidence_registry.md",
        "docs/crop_outcome_decision.md",
        "docs/engine_output_schema.md",
        "notebook/01_AgERA5.ipynb through notebook/10_phase6_exploratory_data_analysis.ipynb",
        "reports/dataset_audit/a8_spatial_grid_cells_audit.csv and reports/dataset_audit/bps_manual_input.csv",
    ])

    paragraph(doc, "End of dossier. The numerical values above are registry values as of 2026-10-12. The intended next action is an independent decision on research scope, followed by a rerun of the canonical Phase 4 pipeline before final closure.")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
