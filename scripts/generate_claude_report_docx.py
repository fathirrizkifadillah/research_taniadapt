"""Convert Claude Report Markdown to formatted Docx.
Location: context/to_claude/TaniAdapt_Laporan_Komprehensif_Desain_Web_dan_Traceability_Riset_ke_Claude.docx
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PROJECT_ROOT = Path(__file__).resolve().parents[1]
MD_PATH = PROJECT_ROOT / "context" / "to_claude" / "TaniAdapt_Laporan_Komprehensif_Desain_Web_dan_Traceability_Riset_ke_Claude.md"
DOCX_OUT = PROJECT_ROOT / "context" / "to_claude" / "TaniAdapt_Laporan_Komprehensif_Desain_Web_dan_Traceability_Riset_ke_Claude.docx"

NAVY = "1B365D"
STEEL = "2C5282"
PALE_BG = "F4F6F9"
BORDER_GRAY = "CCCCCC"


def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def set_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for s in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{s}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), BORDER_GRAY)
        borders.append(b)
    tcPr.append(borders)


def format_cell(cell, text, bold=False, white=False, size=8.5, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    if white:
        run.font.color.rgb = RGBColor(255, 255, 255)
    else:
        run.font.color.rgb = RGBColor(40, 40, 40)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_borders(cell)


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    lines = MD_PATH.read_text(encoding="utf-8").split("\n")

    in_table = False
    table_lines = []

    in_code = False
    code_lines = []

    for line in lines:
        stripped = line.strip()

        # Handle Code Block
        if stripped.startswith("```"):
            if in_code:
                # end code block
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_shading(cell, "F8FAFC")
                tcPr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                left_b = OxmlElement("w:left")
                left_b.set(qn("w:val"), "single")
                left_b.set(qn("w:sz"), "18")
                left_b.set(qn("w:color"), STEEL)
                borders.append(left_b)
                for s in ("top", "bottom", "right"):
                    none_b = OxmlElement(f"w:{s}")
                    none_b.set(qn("w:val"), "none")
                    borders.append(none_b)
                tcPr.append(borders)
                cell.width = Inches(6.5)

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                r = p.add_run("\n".join(code_lines))
                r.font.name = "Consolas"
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(30, 30, 30)
                doc.add_paragraph()
                code_lines = []
                in_code = False
            else:
                in_code = True
                code_lines = []
            continue

        if in_code:
            code_lines.append(line)
            continue

        # Handle Table
        if stripped.startswith("|"):
            in_table = True
            table_lines.append(stripped)
            continue
        else:
            if in_table:
                # process table
                rows = []
                for tl in table_lines:
                    if "---" in tl:
                        continue
                    parts = [p.strip().replace("<br>", "\n") for p in tl.strip("|").split("|")]
                    rows.append(parts)
                if rows:
                    cols = len(rows[0])
                    t = doc.add_table(rows=len(rows), cols=cols)
                    t.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for r_idx, r_data in enumerate(rows):
                        row = t.rows[r_idx]
                        is_hdr = (r_idx == 0)
                        for c_idx, val in enumerate(r_data):
                            if c_idx < cols:
                                cell = row.cells[c_idx]
                                if is_hdr:
                                    set_cell_shading(cell, NAVY)
                                    format_cell(cell, val, bold=True, white=True, size=8.5, center=True)
                                else:
                                    bg = PALE_BG if r_idx % 2 == 1 else "FFFFFF"
                                    if bg != "FFFFFF":
                                        set_cell_shading(cell, bg)
                                    format_cell(cell, val, bold=(c_idx == 0), size=8.2, center=(c_idx == 0 and len(val) < 20))
                    doc.add_paragraph()
                table_lines = []
                in_table = False

        if not stripped:
            continue

        if stripped.startswith("# "):
            h = doc.add_heading(stripped[2:], level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            for r in h.runs:
                r.font.name = "Aptos Display"
                r.font.size = Pt(15)
                r.bold = True
                r.font.color.rgb = RGBColor(27, 54, 93)
        elif stripped.startswith("## "):
            h = doc.add_heading(stripped[3:], level=2)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            for r in h.runs:
                r.font.name = "Aptos"
                r.font.size = Pt(12)
                r.bold = True
                r.font.color.rgb = RGBColor(44, 82, 130)
        elif stripped.startswith("### "):
            h = doc.add_heading(stripped[4:], level=3)
            h.paragraph_format.space_before = Pt(6)
            h.paragraph_format.space_after = Pt(2)
            for r in h.runs:
                r.font.name = "Aptos"
                r.font.size = Pt(10)
                r.bold = True
                r.font.color.rgb = RGBColor(40, 40, 40)
        elif stripped.startswith("---"):
            continue
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            clean_text = stripped.replace("**", "").replace("*", "")
            r = p.add_run(clean_text)
            r.font.name = "Aptos"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(30, 30, 30)

    doc.save(DOCX_OUT)
    print(f"Successfully generated docx at {DOCX_OUT}")


if __name__ == "__main__":
    main()
